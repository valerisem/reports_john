"""Builds the HoM pipeline workbook.

The layout, palette, number formats, conditional formatting and charts are a
faithful rebuild of the supplied template. Everything that was a formula in the
template stays a formula here, so the recipient can still filter, re-sort and
audit the numbers in Excel.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from .model import BLANK, EXISTING, NEW_BUSINESS, ReportData, UNASSIGNED

# -- palette ---------------------------------------------------------------
INDIGO = "250E8D"
PINK = "EE4392"
PURPLE = "8D2A90"
ORANGE = "F8AA57"
ROW_TINT = "F6F5FA"
TOTAL_TINT = "FDE3EF"
GRID = "E2DDF0"
WHITE = "FFFFFF"
BLACK = "000000"
CORAL = "F37A74"

STAGE_TINTS = {
    "Understand Need/Problem": ("ECE6F7", INDIGO),
    "Create Proposal": ("D9CCEB", INDIGO),
    "Present Proposal": ("FDE3EF", INDIGO),
    "Feedback": ("F8B6D3", INDIGO),
    "Negotiation": (PINK, WHITE),
    "IO Sent Out": (PURPLE, WHITE),
}
STATUS_TINTS = {
    NEW_BUSINESS: ("FEEBD6", INDIGO, True),
    EXISTING: ("ECE6F7", PURPLE, False),
}

GBP = "\\£#,##0"
INT = "#,##0"
PCT = "0%"
DATE_SHORT = "dd\\ mmm\\ yy"
DATE_LONG = "dd\\ mmm\\ yyyy"

BOTTOM_RULE = Border(bottom=Side(style="thin", color=GRID))
STALE_DEAL_DAYS = 30
DATA_ROW_HEIGHT = 15.0
TITLE_ROW_HEIGHT = 16.15


def _fill(rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb, bgColor=rgb)


def _font(size=10, bold=False, color=BLACK, underline=None, italic=False) -> Font:
    return Font(name="Arial", size=size, bold=bold, color=color, underline=underline, italic=italic)


def _header_row(ws: Worksheet, row: int, labels: list[tuple[str, str]], height: float = 30) -> None:
    """Write white-on-indigo column headers at (column_letter, label) pairs."""
    ws.row_dimensions[row].height = height
    for column, label in labels:
        cell = ws[f"{column}{row}"]
        cell.value = label
        cell.font = _font(bold=True, color=WHITE)
        cell.fill = _fill(INDIGO)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_body(cell, *, number_format="General", align="general", tint: str | None = None,
                bold=False, color=BLACK, vertical="bottom") -> None:
    cell.font = _font(bold=bold, color=color)
    if tint:
        cell.fill = _fill(tint)
    cell.number_format = number_format
    cell.alignment = Alignment(horizontal=align, vertical=vertical)
    cell.border = BOTTOM_RULE


def _print_setup(ws: Worksheet) -> None:
    """Landscape, scaled to one page wide - matches the template's print setup."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.gridLines = False


# --------------------------------------------------------------------------
def build_workbook(data: ReportData) -> BytesIO:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    deals_ws = wb.create_sheet("Open Deals")
    brands_ws = wb.create_sheet("Brands")
    settings_ws = wb.create_sheet("Settings")

    n_deals = len(data.deals)
    n_brands = len(data.brands)
    deal_last = max(n_deals + 1, 2)
    brand_last = max(n_brands + 1, 2)

    _build_settings(settings_ws, data)
    _build_deals(deals_ws, data, deal_last)
    _build_brands(brands_ws, data, deal_last, brand_last)
    _build_summary(summary, data, deal_last)

    for sheet in (summary, deals_ws, brands_ws, settings_ws):
        _print_setup(sheet)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# -- Settings --------------------------------------------------------------
def _build_settings(ws: Worksheet, data: ReportData) -> None:
    ws.sheet_view.showGridLines = True
    for column, width in {"A": 26, "B": 14, "C": 16}.items():
        ws.column_dimensions[column].width = width

    ws["A1"] = "Settings"
    ws["A1"].font = _font(size=16, bold=True, color=INDIGO)

    ws["A3"] = "Report date"
    ws["B3"] = data.report_date
    ws["B3"].number_format = DATE_LONG

    _header_row(ws, 5, [("A", "Currency"), ("B", "Rate to GBP")])
    for offset, currency in enumerate(("GBP", "USD", "EUR")):
        row = 6 + offset
        ws[f"A{row}"] = currency
        ws[f"B{row}"] = round(data.rates.get(currency, 1.0), 4)
        ws[f"B{row}"].number_format = "0.00" if currency == "GBP" else "0.0000"

    _header_row(ws, 10, [("A", "Stage"), ("B", "Order"), ("C", "Probability")])
    for offset, stage in enumerate(data.stages):
        row = 11 + offset
        ws[f"A{row}"] = stage.name
        ws[f"B{row}"] = stage.order
        ws[f"C{row}"] = stage.probability
        ws[f"C{row}"].number_format = PCT

    note_row = 12 + len(data.stages)
    source = "live FX rates" if data.rates_are_live else "fallback FX rates"
    ws[f"A{note_row}"] = (
        'Probability = Pipedrive stage default. Client status = "Existing client" if the brand '
        f"has at least one won deal in Pipedrive. Values converted using {source}."
    )
    ws[f"A{note_row}"].font = _font(size=9, italic=True)


def _settings_refs(data: ReportData) -> tuple[str, str, str]:
    """(currency range, stage-name range, probability range) on Settings."""
    stage_first, stage_last = 11, 10 + len(data.stages)
    return (
        "Settings!$A$6:$A$8",
        f"Settings!$A${stage_first}:$A${stage_last}",
        f"Settings!$C${stage_first}:$C${stage_last}",
    )


# -- Open Deals ------------------------------------------------------------
DEAL_COLUMNS = [
    ("A", "Brand", 26),
    ("B", "Deal", 44),
    ("C", "Account Owner", 18),
    ("D", "Account Manager", 20),
    ("E", "Client status", 15),
    ("F", "Stage", 22),
    ("G", "Value", 11),
    ("H", "Currency", 9),
    ("I", "Value (£)", 12),
    ("J", "Probability", 11),
    ("K", "Weighted (£)", 13),
    ("L", "Expected close", 14),
    ("M", "Days in stage", 11),
    ("N", "Deal added", 12),
    ("O", "Main contact", 22),
    ("P", "Industry", 18),
    ("Q", "Sub-industry", 24),
    ("R", "Website", 26),
]


def _build_deals(ws: Worksheet, data: ReportData, last_row: int) -> None:
    currency_ref, stage_ref, prob_ref = _settings_refs(data)
    for column, _, width in DEAL_COLUMNS:
        ws.column_dimensions[column].width = width
    _header_row(ws, 1, [(c, label) for c, label, _ in DEAL_COLUMNS])
    ws.freeze_panes = "B2"

    for index, deal in enumerate(data.deals):
        row = index + 2
        values = {
            "A": deal.brand,
            "B": deal.title,
            "C": deal.account_owner,
            "D": deal.account_manager,
            "E": deal.client_status,
            "F": deal.stage,
            "G": deal.value,
            "H": deal.currency,
            "I": f"=G{row}*INDEX(Settings!$B$6:$B$8,MATCH(H{row},{currency_ref},0))",
            "J": f"=INDEX({prob_ref},MATCH(F{row},{stage_ref},0))",
            "K": f"=I{row}*J{row}",
            "L": deal.expected_close,
            "M": (
                f"=Settings!$B$3-DATE({deal.stage_changed.year},{deal.stage_changed.month},{deal.stage_changed.day})"
                if deal.stage_changed
                else None
            ),
            "N": deal.added,
            "O": deal.main_contact,
            "P": deal.industry,
            "Q": deal.sub_industry,
            "R": deal.website,
        }
        formats = {"G": INT, "I": GBP, "J": PCT, "K": GBP, "L": DATE_SHORT, "N": DATE_SHORT}
        centered = {"E", "F", "H", "J", "M"}
        for column, value in values.items():
            cell = ws[f"{column}{row}"]
            cell.value = value
            _style_body(
                cell,
                number_format=formats.get(column, "General"),
                align="center" if column in centered else "general",
                bold=column == "A",
                color=INDIGO if column == "A" else BLACK,
                vertical="center",
            )
        website = ws[f"R{row}"]
        if deal.website:
            website.hyperlink = _as_url(deal.website)
            website.font = _font(color=PURPLE, underline="single")

    ws.auto_filter.ref = f"A1:R{last_row}"
    _apply_row_rules(ws, stage_column="F", status_column="E", weighted_column="K", last_row=last_row)
    ws.conditional_formatting.add(
        f"M2:M{last_row}",
        CellIsRule(operator="greaterThan", formula=[str(STALE_DEAL_DAYS)], font=Font(name="Arial", bold=True, color=CORAL)),
    )
    ws.conditional_formatting.add(
        f"C2:C{last_row}",
        FormulaRule(formula=[f'$C2="{UNASSIGNED}"'], font=Font(name="Arial", bold=True, color=CORAL)),
    )


def _as_url(website: str) -> str:
    website = website.strip()
    if website.startswith(("http://", "https://")):
        return website
    return f"https://{website}"


def _apply_row_rules(ws: Worksheet, *, stage_column: str, status_column: str,
                     weighted_column: str, last_row: int) -> None:
    for stage_name, (bg, fg) in STAGE_TINTS.items():
        ws.conditional_formatting.add(
            f"{stage_column}2:{stage_column}{last_row}",
            FormulaRule(
                formula=[f'${stage_column}2="{stage_name}"'],
                fill=_fill(bg),
                font=Font(name="Arial", bold=True, color=fg),
            ),
        )
    for status, (bg, fg, bold) in STATUS_TINTS.items():
        ws.conditional_formatting.add(
            f"{status_column}2:{status_column}{last_row}",
            FormulaRule(
                formula=[f'${status_column}2="{status}"'],
                fill=_fill(bg),
                font=Font(name="Arial", bold=bold, color=fg),
            ),
        )
    ws.conditional_formatting.add(
        f"{weighted_column}2:{weighted_column}{last_row}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color=PINK, showValue=True),
    )


# -- Brands ----------------------------------------------------------------
BRAND_COLUMNS = [
    ("A", "Brand", 30),
    ("B", "Client status", 15),
    ("C", "Industry", 18),
    ("D", "Sub-industry", 24),
    ("E", "Website", 26),
    ("F", "Open deals", 10),
    ("G", "Pipeline (£)", 14),
    ("H", "Furthest stage", 22),
    ("I", "Account Owner", 18),
    ("J", "Account Manager", 20),
    ("K", "Contacts", 40),
]


def _build_brands(ws: Worksheet, data: ReportData, deal_last: int, last_row: int) -> None:
    _, stage_ref, _ = _settings_refs(data)
    stage_order_ref = stage_ref.replace("$A$", "$B$")
    for column, _, width in BRAND_COLUMNS:
        ws.column_dimensions[column].width = width
    _header_row(ws, 1, [(c, label) for c, label, _ in BRAND_COLUMNS])
    ws.freeze_panes = "B2"

    for index, brand in enumerate(data.brands):
        row = index + 2
        values = {
            "A": brand.name,
            "B": brand.client_status,
            "C": brand.industry,
            "D": brand.sub_industry,
            "E": brand.website,
            "F": f"=COUNTIF('Open Deals'!$A$2:$A${deal_last},A{row})",
            "G": f"=SUMIF('Open Deals'!$A$2:$A${deal_last},A{row},'Open Deals'!$I$2:$I${deal_last})",
            "H": (
                f"=INDEX({stage_ref},SUMPRODUCT(MAX((COUNTIFS('Open Deals'!$A$2:$A${deal_last},A{row},"
                f"'Open Deals'!$F$2:$F${deal_last},{stage_ref})>0)*{stage_order_ref})))"
            ),
            "I": brand.account_owner,
            "J": brand.account_manager,
            "K": ", ".join(brand.contacts),
        }
        for column, value in values.items():
            cell = ws[f"{column}{row}"]
            cell.value = value
            _style_body(
                cell,
                number_format=GBP if column == "G" else "General",
                align="center" if column in {"B", "F"} else "general",
                bold=column == "A",
                color=INDIGO if column == "A" else BLACK,
                vertical="top",
            )
        website = ws[f"E{row}"]
        if brand.website:
            website.hyperlink = _as_url(brand.website)
            website.font = _font(color=PURPLE, underline="single")

    ws.auto_filter.ref = f"A1:K{last_row}"
    _apply_row_rules(ws, stage_column="H", status_column="B", weighted_column="G", last_row=last_row)


# -- Summary ---------------------------------------------------------------
def _kpi(ws: Worksheet, label_cell: str, value_cell: str, span: int, label: str,
         formula: str, fill: str, number_format: str, label_color: str, value_color: str) -> None:
    column = label_cell[0]
    row_label = int(label_cell[1:])
    row_value = int(value_cell[1:])
    start = ws[label_cell].column
    ws.merge_cells(start_row=row_label, start_column=start, end_row=row_label, end_column=start + span - 1)
    ws.merge_cells(start_row=row_value, start_column=start, end_row=row_value, end_column=start + span - 1)

    head = ws[f"{column}{row_label}"]
    head.value = label
    head.font = _font(size=9, bold=True, color=label_color)
    head.fill = _fill(fill)
    head.alignment = Alignment(horizontal="center", vertical="center")

    body = ws[f"{column}{row_value}"]
    body.value = formula
    body.font = _font(size=22, bold=True, color=value_color)
    body.fill = _fill(fill)
    body.number_format = number_format
    body.alignment = Alignment(horizontal="center", vertical="center")
    # Paint the merged tail so the block reads as one solid tile.
    for offset in range(1, span):
        for row in (row_label, row_value):
            ws.cell(row=row, column=start + offset).fill = _fill(fill)


def _table(ws: Worksheet, *, title: str, title_cell: str, columns: list[tuple[str, str]],
           rows: list[list], header_row: int, tint: str | None, total_label: str,
           formats: dict[str, str]) -> int:
    """Render one Summary block; returns the total row number."""
    ws[title_cell] = title
    ws[title_cell].font = _font(size=13, bold=True, color=INDIGO)
    ws.row_dimensions[int(title_cell[1:])].height = TITLE_ROW_HEIGHT
    _header_row(ws, header_row, columns)

    first = header_row + 1
    for offset, values in enumerate(rows):
        row = first + offset
        ws.row_dimensions[row].height = DATA_ROW_HEIGHT
        for (column, _), value in zip(columns, values):
            cell = ws[f"{column}{row}"]
            cell.value = value
            _style_body(
                cell,
                number_format=formats.get(column, "General"),
                align="center" if formats.get(column) in (None, "General", PCT) and column != columns[0][0] else "general",
                tint=tint,
            )
    last = first + len(rows) - 1
    total_row = last + 1
    ws.row_dimensions[total_row].height = DATA_ROW_HEIGHT
    for index, (column, _) in enumerate(columns):
        cell = ws[f"{column}{total_row}"]
        if index == 0:
            cell.value = total_label
        elif column in formats and formats[column] == PCT:
            weighted = columns[-1][0]
            value = columns[-2][0]
            cell.value = f"={weighted}{total_row}/{value}{total_row}"
        else:
            cell.value = f"=SUM({column}{first}:{column}{last})"
        _style_body(
            cell,
            number_format=formats.get(column, "General"),
            align="center" if formats.get(column) in (None, "General", PCT) and index else "general",
            tint=TOTAL_TINT,
            bold=True,
            color=INDIGO,
        )
    return total_row


def _pod_table(ws: Worksheet, *, title_cell: str, header_row: int,
               rows: list[dict]) -> int:
    """Account owners with their pod's account managers indented underneath."""
    ws[title_cell] = "Pods - account owners and their account managers"
    ws[title_cell].font = _font(size=13, bold=True, color=INDIGO)
    ws.row_dimensions[int(title_cell[1:])].height = TITLE_ROW_HEIGHT
    columns = [("B", "Account Owner / Manager"), ("C", "Deals"), ("D", "Value (£)"), ("E", "Weighted (£)")]
    _header_row(ws, header_row, columns)

    first = header_row + 1
    for offset, row in enumerate(rows):
        excel_row = first + offset
        ws.row_dimensions[excel_row].height = DATA_ROW_HEIGHT
        is_owner = row["manager"] is None
        label = row["owner"] if is_owner else f"    {row['manager']}"
        values = [label, row["deals"], row["value_gbp"], row["weighted_gbp"]]
        for (column, _), value in zip(columns, values):
            cell = ws[f"{column}{excel_row}"]
            cell.value = value
            _style_body(
                cell,
                number_format={"C": "General", "D": GBP, "E": GBP}.get(column, "General"),
                align="center" if column == "C" else "general",
                tint=ROW_TINT if is_owner else None,
                bold=is_owner,
                color=INDIGO if is_owner else BLACK,
            )
    return first + len(rows) - 1


def _build_summary(ws: Worksheet, data: ReportData, deal_last: int) -> None:
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 24, "C": 11, "D": 13, "E": 15, "G": 3, "H": 24, "I": 9, "J": 15}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.row_dimensions[1].height = 36
    ws["B1"] = "Sales Pipeline Report"
    ws["B1"].font = _font(size=22, bold=True, color=INDIGO)
    ws["B2"] = '="All open deals as of "&TEXT(Settings!B3,"dd mmm yyyy")&"  |  values in GBP"'
    ws["B2"].font = _font(size=11, color=PURPLE)
    ws.row_dimensions[2].height = DATA_ROW_HEIGHT

    ws.row_dimensions[4].height = 21.75
    ws.row_dimensions[5].height = 42
    _kpi(ws, "B4", "B5", 2, "OPEN DEALS", f"=COUNTA('Open Deals'!$B$2:$B${deal_last})",
         INDIGO, "0", WHITE, WHITE)
    _kpi(ws, "D4", "D5", 3, "PIPELINE VALUE", f"=SUM('Open Deals'!$I$2:$I${deal_last})",
         PINK, GBP, WHITE, WHITE)
    _kpi(ws, "H4", "H5", 2, "WEIGHTED VALUE", f"=SUM('Open Deals'!$K$2:$K${deal_last})",
         PURPLE, GBP, WHITE, WHITE)
    _kpi(ws, "J4", "J5", 2, "NEW BUSINESS (£)",
         f"=SUMIF('Open Deals'!$E$2:$E${deal_last},\"{NEW_BUSINESS}\",'Open Deals'!$I$2:$I${deal_last})",
         ORANGE, GBP, INDIGO, INDIGO)

    _, stage_ref, prob_ref = _settings_refs(data)
    deals_col = lambda letter: f"'Open Deals'!${letter}$2:${letter}${deal_last}"  # noqa: E731

    def count(match_col: str, key_cell: str) -> str:
        return f"=COUNTIF({deals_col(match_col)},{key_cell})"

    def total(match_col: str, key_cell: str, sum_col: str) -> str:
        return f"=SUMIF({deals_col(match_col)},{key_cell},{deals_col(sum_col)})"

    # Block 1 - stages (left) and account owners (right).
    stage_rows = [
        [
            stage.name,
            count("F", f"B{10 + i}"),
            f"=INDEX({prob_ref},MATCH(B{10 + i},{stage_ref},0))",
            total("F", f"B{10 + i}", "I"),
            total("F", f"B{10 + i}", "K"),
        ]
        for i, stage in enumerate(data.stages)
    ]
    stage_total = _table(
        ws, title="Pipeline by stage", title_cell="B8",
        columns=[("B", "Stage"), ("C", "Deals"), ("D", "Probability %"), ("E", "Value (£)"), ("F", "Weighted (£)")],
        rows=stage_rows, header_row=9, tint=ROW_TINT, total_label="Total",
        formats={"C": "General", "D": PCT, "E": GBP, "F": GBP},
    )
    owner_rows = [
        [name, count("C", f"H{10 + i}"), total("C", f"H{10 + i}", "I"), total("C", f"H{10 + i}", "K")]
        for i, name in enumerate(data.account_owners)
    ]
    owner_total = _table(
        ws, title="Pipeline by Account Owner", title_cell="H8",
        columns=[("H", "Account Owner"), ("I", "Deals"), ("J", "Value (£)"), ("K", "Weighted (£)")],
        rows=owner_rows, header_row=9, tint=ROW_TINT, total_label="Total",
        formats={"I": "General", "J": GBP, "K": GBP},
    )

    # Block 2 - account managers (left) and industries (right).
    header2 = max(stage_total, owner_total) + 3
    manager_total = _pod_table(
        ws, title_cell=f"B{header2 - 1}", header_row=header2, rows=data.by_pod()
    )
    industry_rows = [
        [name, count("P", f"H{header2 + 1 + i}"), total("P", f"H{header2 + 1 + i}", "I"),
         total("P", f"H{header2 + 1 + i}", "K")]
        for i, name in enumerate(data.industries)
    ]
    industry_total = _table(
        ws, title="Pipeline by industry", title_cell=f"H{header2 - 1}",
        columns=[("H", "Industry"), ("I", "Deals"), ("J", "Value (£)"), ("K", "Weighted (£)")],
        rows=industry_rows, header_row=header2, tint=None, total_label="Total",
        formats={"I": "General", "J": GBP, "K": GBP},
    )

    charts_row = max(manager_total, industry_total) + 2
    ws[f"B{charts_row}"] = "Charts"
    ws[f"B{charts_row}"].font = _font(size=13, bold=True, color=INDIGO)
    ws.row_dimensions[charts_row].height = TITLE_ROW_HEIGHT

    stage_chart = _bar_chart(
        ws, title="Pipeline value by stage (£)", colour=PINK,
        cats=Reference(ws, min_col=2, min_row=10, max_row=stage_total - 1),
        vals=Reference(ws, min_col=5, min_row=10, max_row=stage_total - 1),
    )
    owner_chart = _bar_chart(
        ws, title="Pipeline value by Account Owner (£)", colour=PURPLE,
        cats=Reference(ws, min_col=8, min_row=10, max_row=owner_total - 1),
        vals=Reference(ws, min_col=10, min_row=10, max_row=owner_total - 1),
    )
    ws.add_chart(stage_chart, f"B{charts_row + 1}")
    ws.add_chart(owner_chart, f"H{charts_row + 1}")


def _bar_chart(ws: Worksheet, *, title: str, colour: str, cats: Reference, vals: Reference) -> BarChart:
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "clustered"
    chart.title = title
    chart.legend = None
    chart.add_data(vals, titles_from_data=False)
    chart.set_categories(cats)
    chart.width = 15.7
    chart.height = 8.1
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.numFmt = GBP
    chart.dLbls.showSerName = False
    chart.dLbls.showCatName = False
    chart.dLbls.showLegendKey = False
    series = chart.series[0]
    series.graphicalProperties.solidFill = colour
    series.graphicalProperties.line = LineProperties(solidFill=colour)
    chart.gapWidth = 50
    # In openpyxl x_axis is always the category axis, even for horizontal bars.
    # Bar charts draw the first category at the bottom, so reverse the category
    # axis to read top-down in pipeline order, and drop the value axis because
    # every bar already carries its own figure.
    chart.x_axis.scaling.orientation = "maxMin"
    chart.x_axis.delete = False
    chart.y_axis.delete = True
    chart.y_axis.majorGridlines = None
    return chart
