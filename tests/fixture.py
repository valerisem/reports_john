"""Builds Pipedrive-shaped payloads from the supplied reference workbook.

This lets the whole pipeline (model -> Excel -> email) be exercised against the
129 real deals in the template, without needing API credentials.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl

REFERENCE = Path(__file__).resolve().parent.parent / "assets" / "template_reference.xlsx"

STAGES = [
    {"id": 2, "name": "Understand Need/Problem", "order_nr": 0, "deal_probability": 45, "pipeline_id": 1},
    {"id": 3, "name": "Create Proposal", "order_nr": 1, "deal_probability": 50, "pipeline_id": 1},
    {"id": 4, "name": "Present Proposal", "order_nr": 2, "deal_probability": 55, "pipeline_id": 1},
    {"id": 15, "name": "Feedback", "order_nr": 3, "deal_probability": 60, "pipeline_id": 1},
    {"id": 6, "name": "Negotiation", "order_nr": 4, "deal_probability": 75, "pipeline_id": 1},
    {"id": 7, "name": "IO Sent Out", "order_nr": 5, "deal_probability": 85, "pipeline_id": 1},
]

AM_KEY = "am_custom_key"
INDUSTRY_KEY = "industry_key"
SUB_INDUSTRY_KEY = "sub_industry_key"
FIELD_KEYS = {
    "deal_account_manager": AM_KEY,
    "org_industry": INDUSTRY_KEY,
    "org_sub_industry": SUB_INDUSTRY_KEY,
}


def _iso(value) -> str | None:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return None


def load() -> dict:
    wb = openpyxl.load_workbook(REFERENCE, data_only=False)
    deals_ws, brands_ws = wb["Open Deals"], wb["Brands"]
    stage_by_name = {s["name"]: s for s in STAGES}

    # Brand metadata keyed by name, straight off the Brands sheet.
    brand_meta: dict[str, dict] = {}
    for row in brands_ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        brand_meta[row[0]] = {
            "client_status": row[1],
            "industry": row[2],
            "sub_industry": row[3],
            "website": row[4],
            "contacts": [c.strip() for c in (row[10] or "").split(",") if c.strip()],
        }

    users: dict[int, str] = {}
    def user_id(name: str | None) -> int | None:
        if not name or name in ("-", ""):
            return None
        for uid, existing in users.items():
            if existing == name:
                return uid
        uid = 1000 + len(users)
        users[uid] = name
        return uid

    orgs: dict[int, dict] = {}
    org_ids: dict[str, int] = {}
    persons: dict[int, dict] = {}
    org_contacts: dict[int, list[dict]] = {}
    deals: list[dict] = []
    won_org_ids: set[int] = set()

    def org_id_for(brand: str) -> int:
        if brand in org_ids:
            return org_ids[brand]
        oid = 500 + len(org_ids)
        org_ids[brand] = oid
        meta = brand_meta.get(brand, {})
        orgs[oid] = {
            "id": oid,
            "name": brand,
            "website": meta.get("website"),
            "custom_fields": {
                INDUSTRY_KEY: meta.get("industry"),
                SUB_INDUSTRY_KEY: meta.get("sub_industry"),
            },
        }
        org_contacts[oid] = []
        for contact in meta.get("contacts", []):
            pid = 9000 + len(persons)
            persons[pid] = {"id": pid, "name": contact, "org_id": oid}
            org_contacts[oid].append(persons[pid])
        if meta.get("client_status") == "Existing client":
            won_org_ids.add(oid)
        return oid

    for index, row in enumerate(deals_ws.iter_rows(min_row=2, values_only=True)):
        brand = row[0]
        if not brand:
            continue
        oid = org_id_for(brand)
        stage = stage_by_name[row[5]]
        contact_name = row[14]
        person_id = None
        if contact_name:
            for pid, person in persons.items():
                if person["org_id"] == oid and person["name"] == contact_name:
                    person_id = pid
                    break
            if person_id is None:
                person_id = 9000 + len(persons)
                persons[person_id] = {"id": person_id, "name": contact_name, "org_id": oid}
                org_contacts[oid].append(persons[person_id])

        # "Days in stage" in the template is a formula: Settings!B3 - DATE(...).
        stage_changed = None
        formula = row[12]
        if isinstance(formula, str) and "DATE(" in formula:
            args = formula.split("DATE(")[1].rstrip(")").split(",")
            stage_changed = dt.date(int(args[0]), int(args[1]), int(args[2]))

        deals.append(
            {
                "id": 10_000 + index,
                "title": row[1],
                "org_id": oid,
                "person_id": person_id,
                "owner_id": user_id(row[2]),
                "stage_id": stage["id"],
                "pipeline_id": 1,
                "status": "open",
                "value": row[6],
                "currency": row[7],
                "expected_close_date": _iso(row[11]),
                "add_time": (_iso(row[13]) or "") + "T00:00:00Z" if row[13] else None,
                "stage_change_time": (stage_changed.isoformat() + "T00:00:00Z") if stage_changed else None,
                "custom_fields": {AM_KEY: user_id(row[3])},
            }
        )

    return {
        "stages_payload": STAGES,
        "deals_payload": deals,
        "orgs": orgs,
        "persons": persons,
        "org_contacts": org_contacts,
        "users": users,
        "won_org_ids": won_org_ids,
        "field_keys": FIELD_KEYS,
        "rates": {"GBP": 1.0, "USD": 0.74, "EUR": 0.85},
        "rates_are_live": False,
        "report_date": dt.date(2026, 9, 17),
    }
