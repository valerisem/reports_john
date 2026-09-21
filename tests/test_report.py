"""End-to-end checks against the 129 real deals in the reference workbook."""
from __future__ import annotations

import io

import openpyxl
import pytest

from app.email_html import _leaderboard, _stage_segments, render_email
from app.excel import build_workbook
from app.formatting import bar_percent, compact_gbp, first_name, percent
from app.model import EXISTING, NEW_BUSINESS, build_report
from tests import fixture


@pytest.fixture(scope="module")
def data():
    return build_report(**fixture.load())


@pytest.fixture(scope="module")
def workbook(data):
    return openpyxl.load_workbook(io.BytesIO(build_workbook(data).getvalue()))


# -- the numbers -----------------------------------------------------------
def test_headline_numbers_match_the_reference_report(data):
    assert data.open_deal_count == 129
    # 103, not the reference workbook's 104: "Olymptrade" and
    # "Maree/Olymptrade" share a website and are now one brand.
    assert len(data.brands) == 103
    assert data.new_brand_count == 53
    assert round(data.pipeline_gbp) == 7_430_747
    assert round(data.weighted_gbp) == 4_352_052


def test_stage_totals_match_the_reference_report(data):
    expected = {
        "Understand Need/Problem": (33, 1_586_687),
        "Create Proposal": (21, 1_409_150),
        "Present Proposal": (12, 991_590),
        "Feedback": (28, 1_803_900),
        "Negotiation": (23, 877_540),
        "IO Sent Out": (12, 761_880),
    }
    actual = {row["stage"]: (row["deals"], round(row["value_gbp"])) for row in data.by_stage()}
    assert actual == expected


def test_owner_totals_match_the_reference_report(data):
    expected = {
        "Valeriia Mukhai": (50, 1_918_792.5),
        "Ritchie Boubouli": (40, 1_242_859.15),
        "Carrick Klopper": (39, 1_190_400.0),
    }
    actual = {row["name"]: (row["deals"], row["weighted_gbp"]) for row in data.by_owner()}
    assert actual.keys() == expected.keys()
    for name, (deals, weighted) in expected.items():
        assert actual[name][0] == deals
        assert actual[name][1] == pytest.approx(weighted, abs=0.01)


def test_deals_are_sorted_by_stage_then_value(data):
    keys = [(-d.stage_order, -d.value_gbp) for d in data.deals]
    assert keys == sorted(keys)


def test_every_deal_carries_a_client_status(data):
    assert {d.client_status for d in data.deals} <= {EXISTING, NEW_BUSINESS}
    assert all(d.client_status for d in data.deals)


def test_weighted_value_is_value_times_stage_probability(data):
    for deal in data.deals:
        assert deal.weighted_gbp == pytest.approx(deal.value_gbp * deal.probability)


def test_brand_pipeline_equals_the_sum_of_its_deals(data):
    for brand in data.brands:
        deals = [d for d in data.deals if d.brand == brand.name]
        assert brand.open_deals == len(deals)
        assert brand.pipeline_gbp == pytest.approx(sum(d.value_gbp for d in deals))


# -- formatting ------------------------------------------------------------
@pytest.mark.parametrize(
    "value,expected",
    [
        (4_352_052, "£4.35m"),
        (1_803_900, "£1.80m"),
        (991_590, "£992k"),
        (388_500, "£388k"),
        (74_000, "£74k"),
        (450, "£450"),
        (0, "£0"),
    ],
)
def test_compact_gbp_matches_the_reference_email(value, expected):
    assert compact_gbp(value) == expected


def test_stage_bars_reproduce_the_reference_widths():
    # Values and widths taken straight from the supplied email.
    values = [1_586_687, 1_409_150, 991_590, 1_803_900, 877_540, 761_880]
    maximum = max(values)
    assert [bar_percent(v, maximum, 94) for v in values] == [83, 73, 52, 94, 46, 40]


def test_owner_bars_reproduce_the_reference_widths():
    values = [1_918_793, 1_242_859, 1_190_400]
    assert [bar_percent(v, max(values)) for v in values] == [100, 65, 62]


def test_first_name_shortens_owner_labels():
    assert first_name("Valeriia Mukhai") == "Valeriia"
    assert first_name("") == ""


# -- workbook --------------------------------------------------------------
def test_workbook_has_the_template_sheets(workbook):
    assert workbook.sheetnames == ["Summary", "Open Deals", "Brands", "Settings"]


def test_workbook_row_counts_match_the_data(workbook, data):
    assert workbook["Open Deals"].max_row == len(data.deals) + 1
    assert workbook["Brands"].max_row == len(data.brands) + 1


def test_open_deals_headers_match_the_template(workbook):
    headers = [c.value for c in workbook["Open Deals"][1]]
    assert headers == [
        "Brand", "Deal", "Account Owner", "Account Manager", "Client status", "Stage",
        "Value", "Currency", "Value (£)", "Probability", "Weighted (£)", "Expected close",
        "Days in stage", "Deal added", "Main contact", "Industry", "Sub-industry", "Website",
    ]


def test_brands_headers_match_the_template(workbook):
    headers = [c.value for c in workbook["Brands"][1]]
    assert headers == [
        "Brand", "Client status", "Industry", "Sub-industry", "Website", "Open deals",
        "Pipeline (£)", "Furthest stage", "Account Owner", "Account Manager", "Contacts",
        "Delivered campaigns", "Delivered revenue (£)", "Gross profit (£)", "Gross margin",
        "Margin basis",
    ]


def test_value_columns_stay_formulas_so_the_sheet_stays_auditable(workbook):
    ws = workbook["Open Deals"]
    assert str(ws["I2"].value).startswith("=")
    assert str(ws["J2"].value).startswith("=")
    assert ws["K2"].value == "=I2*K2".replace("K2", "J2")


def test_summary_kpis_reference_the_full_deal_range(workbook, data):
    last = len(data.deals) + 1
    assert workbook["Summary"]["B5"].value == f"=COUNTA('Open Deals'!$B$2:$B${last})"
    assert workbook["Summary"]["D5"].value == f"=SUM('Open Deals'!$I$2:$I${last})"


def test_settings_sheet_carries_rates_and_stage_probabilities(workbook, data):
    ws = workbook["Settings"]
    assert ws["B3"].value.date() == data.report_date
    assert {ws[f"A{r}"].value for r in (6, 7, 8)} == {"GBP", "USD", "EUR"}
    stage_names = [ws[f"A{11 + i}"].value for i in range(len(data.stages))]
    assert stage_names == [s.name for s in data.stages]


def test_charts_are_attached_to_the_summary(workbook):
    assert len(workbook["Summary"]._charts) == 2


def test_workbook_is_landscape_fit_to_width(workbook):
    for name in workbook.sheetnames:
        assert workbook[name].page_setup.orientation == "landscape"


# -- email -----------------------------------------------------------------
@pytest.fixture(scope="module")
def html(data):
    return render_email(data, title="Sales Pipeline Update",
                        greeting_name="John", sender_name="Valeria")


def test_email_greets_the_recipient_and_signs_off(html):
    assert "Hi John, here is the latest pipeline update." in html
    assert "Valeria" in html


def test_email_shows_the_headline_stats(html, data):
    assert f">{data.open_deal_count:,}<" in html
    assert compact_gbp(data.weighted_gbp) in html
    # The third tile is the new/retained split by value, not a brand count:
    # a raw count of new brands goes stale and repeats week after week.
    assert percent(data.new_business_share) in html
    assert "New business" in html


def test_email_lists_ten_brands_per_column(html, data):
    assert "Top 10 retained clients" in html
    assert "New this week" in html  # wording follows the cadence
    # Without history nothing counts as new, so that column is empty.
    assert data.brands_new_this_report == []
    assert "No new brands since the last update." in html
    # The fallback fills the column so it is never near-empty.
    assert "Top new business in the pipeline" in html
    assert html.count("POC:") == 5


def test_email_counts_the_brands_not_shown(html, data):
    shown = min(10, sum(1 for b in data.brands if b.client_status == EXISTING))
    shown += len(data.brands_new_this_report[:10])
    if len(data.brands_new_this_report) < 3:
        shown += 5      # the fallback list
    assert f"{len(data.brands) - shown} other brands, retained and new are in the pipeline" in html


def test_email_uses_friendly_stage_labels(html):
    for label in ("Discovery", "Preparing proposal", "Proposal presented",
                  "Awaiting feedback", "Negotiating", "Contract sent"):
        assert label in html


def test_test_mode_banner_names_the_real_recipients(data):
    marked = render_email(
        data, title="T", greeting_name="John", sender_name="V",
        test_banner={"to": "john@example.com", "cc": "ops@example.com"},
    )
    assert "TEST MODE" in marked
    assert "john@example.com" in marked and "ops@example.com" in marked
    assert "TEST MODE" not in render_email(data, title="T", greeting_name="John", sender_name="V")


# -- the leaderboard and the stacked stage bars ----------------------------
def test_stage_owner_split_adds_up_to_each_stage_total(data):
    split = data.by_stage_owner()
    for row in data.by_stage():
        if not row["deals"]:
            continue
        assert sum(split[row["stage"]].values()) == pytest.approx(row["value_gbp"])


def test_leaderboard_runs_winner_first(data):
    entries = _leaderboard(data.by_owner(), data.by_pod(), "https://example.test/{name}.png")
    assert [e["rank"] for e in entries] == [1, 2, 3]
    assert entries[0]["name"] == "Valeriia"
    # the tallest column belongs to the winner, and the order matches the bars
    assert entries[0]["height"] == max(e["height"] for e in entries)


def test_pod_splits_into_managers_plus_the_leads_own_share(data):
    entries = _leaderboard(data.by_owner(), data.by_pod(), "")
    pods = {r["owner"]: r for r in data.by_pod() if r["manager"] is None}
    for entry in entries:
        managed = sum(
            r["weighted_gbp"] for r in data.by_pod()
            if r["owner"] == entry["owner"] and r["manager"]
        )
        total = pods[entry["owner"]]["weighted_gbp"]
        # what the lead handles plus what their managers handle is the whole pod
        assert compact_gbp(total - managed) == entry["direct"]
        assert entry["direct_percent"] == round((total - managed) / total * 100)


def test_photos_are_named_after_plain_first_names(data):
    entries = _leaderboard(data.by_owner(), data.by_pod(), "https://example.test/{name}.png")
    assert entries[0]["photo"] == "https://example.test/valeriia.png"
    assert all(e["photo"] for e in entries)
    # a hyphenated first name still resolves to the plain one: emma.png
    managers = [m for e in entries for m in e["managers"]]
    assert managers, "the reference data should place managers inside pods"
    assert all(m["photo"].endswith(".png") for m in managers)
    assert all("-" not in m["photo"].rsplit("/", 1)[-1] for m in managers)
    blank = _leaderboard(data.by_owner(), data.by_pod(), "")
    assert all(e["photo"] == "" for e in blank)


def test_stage_segments_keep_one_order_and_fill_the_bar(data):
    entries = _leaderboard(data.by_owner(), data.by_pod(), "")
    split = data.by_stage_owner()
    orders = []
    for row in data.by_stage():
        if not row["deals"]:
            continue
        percent = bar_percent(row["value_gbp"], max(r["value_gbp"] for r in data.by_stage()), 94)
        segments = _stage_segments(row["stage"], row["value_gbp"], percent, split, entries)
        assert sum(s["percent"] for s in segments) == percent
        orders.append([s["colour"] for s in segments])
    # every bar stacks the owners in the same order, so the bars compare
    assert len(set(tuple(o) for o in orders)) == 1


def test_email_renders_the_leaderboard_ahead_of_the_stage_chart(data):
    html = render_email(data, title="T", greeting_name="John", sender_name="V")
    assert html.index("Weighted Pipeline Leaderboard") < html.index("Pipeline by stage")
    assert "Weighted by account owner" not in html
    assert "/valeriia.png" in html
    assert "self-managed" in html


# -- brand identity in the report ------------------------------------------
def test_duplicate_organisation_records_collapse_into_one_brand(data):
    """Records sharing a website are one brand, so the figures are not split."""
    names = [b.name for b in data.brands]
    assert len(names) == len(set(names))
    assert "Maree/Olymptrade" not in names
    olymptrade = next(b for b in data.brands if b.name == "Olymptrade")
    assert olymptrade.is_duplicated is True


def test_every_brand_carries_a_stable_key(data):
    keys = [b.brand_key for b in data.brands]
    assert all(keys)
    assert len(keys) == len(set(keys))


def test_nothing_is_flagged_new_without_history(data):
    """A Supabase outage should quieten the email, not re-announce everything."""
    assert data.history.loaded is False
    assert data.brands_new_this_report == []


def test_new_business_share_is_value_based(data):
    assert data.new_business_share == pytest.approx(
        data.new_business_gbp / data.pipeline_gbp
    )
    assert 0.0 < data.new_business_share < 1.0


def test_the_loaded_history_reaches_the_report():
    """Regression: the history was fetched, then dropped on the floor, so
    every brand looked un-announced no matter what Supabase held."""
    from app.brand_history import BrandHistory

    payload = dict(fixture.load())
    known = BrandHistory(reported_keys={"runwayml.com"}, loaded=True)
    built = build_report(**payload, history=known)
    assert built.history.loaded is True
    assert built.history.reported_keys == {"runwayml.com"}


# -- quiet periods ------------------------------------------------------
def _with_history(reported: set[str]):
    from app.brand_history import BrandHistory

    return build_report(**dict(fixture.load()), history=BrandHistory(reported_keys=reported, loaded=True))


def test_a_quiet_period_falls_back_to_the_biggest_new_business():
    """Nothing new: the column still carries the top new business rather than
    showing John an empty box."""
    data = _with_history({b.brand_key for b in build_report(**fixture.load()).brands})
    assert data.brands_new_this_report == []
    html = render_email(data, title="T", greeting_name="John", sender_name="V")
    assert "No new brands since the last update." in html
    assert "Top new business in the pipeline" in html
    assert html.count("POC:") == 5


def test_a_busy_period_shows_only_the_genuinely_new():
    """Three or more new brands is enough on its own; no fallback."""
    data = _with_history(set())
    assert len(data.brands_new_this_report) >= 3
    html = render_email(data, title="T", greeting_name="John", sender_name="V")
    assert "Top new business in the pipeline" not in html


def test_the_fallback_never_repeats_a_brand_already_listed():
    all_brands = build_report(**fixture.load()).brands
    new_business = sorted(
        (b for b in all_brands if b.client_status == NEW_BUSINESS),
        key=lambda b: -b.weighted_gbp,
    )
    # Everything reported except the single biggest new-business brand.
    data = _with_history({b.brand_key for b in all_brands} - {new_business[0].brand_key})
    assert len(data.brands_new_this_report) == 1
    html = render_email(data, title="T", greeting_name="John", sender_name="V")
    assert "Top new business in the pipeline" in html
    assert html.count(f">{new_business[0].name}<") == 1


# -- deals with no organisation linked ------------------------------------
def _payload_with_unlinked_deal(title: str, value: float = 60_000.0) -> dict:
    """The reference payload plus one deal whose org_id is None."""
    payload = fixture.load()
    template = payload["deals_payload"][0]
    payload["deals_payload"] = payload["deals_payload"] + [
        {
            **template,
            "id": 999_999,
            "title": title,
            "org_id": None,
            "person_id": None,
            "value": value,
            "currency": "GBP",
        }
    ]
    return payload


def test_a_deal_with_no_organisation_still_counts(data):
    """Dropping it would silently under-report the stage and pipeline totals."""
    payload = _payload_with_unlinked_deal("Pemberton Tea Rooms x TikTok Campaign")
    report = build_report(**payload)

    assert report.open_deal_count == data.open_deal_count + 1
    assert round(report.pipeline_gbp) == round(data.pipeline_gbp) + 60_000

    kept = report.deals_missing_organisation[0]
    before = next(s for s in data.by_stage() if s["stage"] == kept.stage)
    after = next(s for s in report.by_stage() if s["stage"] == kept.stage)
    assert after["deals"] == before["deals"] + 1


def test_an_unlinked_deal_takes_its_brand_from_the_title():
    report = build_report(**_payload_with_unlinked_deal("Pemberton Tea Rooms x TikTok Campaign"))
    kept = report.deals_missing_organisation
    assert [d.brand for d in kept] == ["Pemberton Tea Rooms"]

    brand = next(b for b in report.brands if b.name == "Pemberton Tea Rooms")
    # No brand key, so it can never be announced to John as new or written to
    # the history table on the strength of a deal title alone.
    assert brand.brand_key == ""
    assert brand.is_new_this_report is False
    assert brand.name not in {b["brand_name"] for b in report.brand_records()}


def test_an_unlinked_deal_rejoins_a_brand_it_names_exactly(data):
    known = data.brands[0].name
    report = build_report(**_payload_with_unlinked_deal(f"{known} x Winter Campaign"))

    assert len(report.brands) == len(data.brands)  # no duplicate brand row
    brand = next(b for b in report.brands if b.name == known)
    assert brand.open_deals == data.brands[0].open_deals + 1
    assert brand.client_status == data.brands[0].client_status
    assert brand.brand_key == data.brands[0].brand_key


def test_an_unlinked_deal_with_no_usable_title_is_dropped(data):
    report = build_report(**_payload_with_unlinked_deal(""))
    assert report.open_deal_count == data.open_deal_count


# -- won year to date ------------------------------------------------------
import datetime as _dt

from app.config import Settings


def _won_payload(owner_ids: list[int]) -> list[dict]:
    won = []
    for i, oid in enumerate(owner_ids):
        won.append({
            "id": 900_000 + i, "title": f"Won deal {i}", "org_id": None, "owner_id": oid,
            "value": 50_000 * (i + 1), "currency": "GBP",
            "won_time": f"2026-06-0{i + 1}T10:00:00Z",
        })
    return won


def _ytd(data_payload: dict, won: list[dict]):
    return build_report(**data_payload, won_deals_payload=won,
                        financial_year_start=_dt.date(2026, 4, 1))


def test_won_ytd_totals_and_owner_split():
    payload = fixture.load()
    owner_ids = list(payload["users"])[:3]
    report = _ytd(payload, _won_payload(owner_ids))

    assert len(report.won_ytd) == 3
    assert round(report.won_ytd_gbp) == 50_000 + 100_000 + 150_000
    rows = {r["name"]: r for r in report.ytd_by_owner()}
    assert sum(r["deals"] for r in rows.values()) == 3
    # Biggest first, and every pipeline owner is listed even on a blank year.
    values = [r["value_gbp"] for r in report.ytd_by_owner()]
    assert values == sorted(values, reverse=True)
    assert set(report.account_owners) <= set(rows)


def test_won_before_the_financial_year_is_excluded():
    payload = fixture.load()
    owner_id = list(payload["users"])[0]
    won = _won_payload([owner_id]) + [{
        "id": 1, "title": "Last year", "org_id": None, "owner_id": owner_id,
        "value": 999_999, "currency": "GBP", "won_time": "2025-06-01T10:00:00Z",
    }]
    report = _ytd(payload, won)
    assert len(report.won_ytd) == 1
    assert round(report.won_ytd_gbp) == 50_000


def test_won_ytd_is_converted_to_sterling():
    payload = fixture.load()
    owner_id = list(payload["users"])[0]
    won = [{"id": 1, "title": "USD win", "org_id": None, "owner_id": owner_id,
            "value": 100_000, "currency": "USD", "won_time": "2026-06-01T10:00:00Z"}]
    report = _ytd(payload, won)
    assert round(report.won_ytd_gbp) == round(100_000 * payload["rates"]["USD"])


def test_email_hides_the_ytd_chart_when_nothing_is_won(data):
    html = render_email(data, title="t", greeting_name="John", sender_name="Valeria")
    assert "Won year to date" not in html


def test_the_email_has_no_separate_year_to_date_chart():
    """Year to date belongs beside each person, not in a chart of its own."""
    payload = fixture.load()
    report = _ytd(payload, _won_payload(list(payload["users"])[:2]))
    html = render_email(report, title="t", greeting_name="John", sender_name="Valeria")
    assert "Won year to date" not in html
    # The figure itself still reaches the reader, on the podium.
    assert "won YTD" in html


def test_financial_year_start_follows_the_configured_month():
    april = Settings(pipedrive_api_token="x")
    assert april.financial_year_start(_dt.date(2026, 9, 20)) == _dt.date(2026, 4, 1)
    assert april.financial_year_start(_dt.date(2026, 3, 31)) == _dt.date(2025, 4, 1)
    calendar = Settings(pipedrive_api_token="x", financial_year_start_month=1)
    assert calendar.financial_year_start(_dt.date(2026, 9, 20)) == _dt.date(2026, 1, 1)


# -- campaign profitability ------------------------------------------------
from app.campaign_finance import CampaignFinanceSet, build_finance


def test_bookings_prefixed_board_or_pay_are_ignored():
    """Those rows are creators listed on a board, never paid - no fee at all."""
    finance = build_finance(
        [{"pd_deal_id": 5, "client_name": "X", "stage": "Delivered",
          "budget_gbp": 100_000, "current_spend_gbp": 40_000}],
        [{"campaign_number": "5", "fee_gbp": 30_000},
         {"campaign_number": "board:99", "fee_gbp": None},
         {"campaign_number": "pay:123", "fee_gbp": 5_000}],
    )
    campaign = finance.by_deal[5]
    assert campaign.influencer_cost_gbp == 30_000
    assert campaign.costed_from_payments is True


def test_cost_falls_back_to_board_spend_when_no_payments_exist():
    finance = build_finance(
        [{"pd_deal_id": 7, "client_name": "X", "stage": "Delivered",
          "budget_gbp": 50_000, "current_spend_gbp": 20_000}],
        [],
    )
    assert finance.by_deal[7].influencer_cost_gbp == 20_000
    assert finance.by_deal[7].costed_from_payments is False


def test_margin_includes_paid_media_and_brand_uplift():
    finance = build_finance(
        [{"pd_deal_id": 1, "client_name": "X", "stage": "Delivered", "budget_gbp": 100_000,
          "paid_media_spend": 10_000, "brand_uplift_spend": 5_000}],
        [{"campaign_number": "1", "fee_gbp": 35_000}],
    )
    rolled = finance.roll_up({1})
    assert rolled.cost_gbp == 50_000
    assert rolled.gross_profit_gbp == 50_000
    assert rolled.margin == pytest.approx(0.5)


def test_campaigns_still_running_are_left_out_of_margin():
    finance = build_finance(
        [{"pd_deal_id": 1, "client_name": "X", "stage": "WIP",
          "budget_gbp": 100_000, "current_spend_gbp": 1_000}],
        [{"campaign_number": "1", "fee_gbp": 1_000}],
    )
    assert finance.roll_up({1}).margin is None


def test_duplicate_client_records_total_as_one_brand():
    """Opera, Opera Ltd and Opera Browser are one client, not three."""
    finance = build_finance(
        [{"pd_deal_id": 11, "client_name": "Opera", "stage": "Delivered", "budget_gbp": 100_000},
         {"pd_deal_id": 12, "client_name": "Opera Ltd", "stage": "Delivered", "budget_gbp": 100_000},
         {"pd_deal_id": 13, "client_name": "Opera Browser", "stage": "Delivered", "budget_gbp": 100_000}],
        [{"campaign_number": "11", "fee_gbp": 70_000},
         {"campaign_number": "12", "fee_gbp": 50_000},
         {"campaign_number": "13", "fee_gbp": 30_000}],
    )
    # All three deals sit on organisations that resolve to a single brand.
    rolled = finance.roll_up({11, 12, 13})
    assert rolled.campaigns == 3
    assert rolled.revenue_gbp == 300_000
    assert rolled.margin == pytest.approx(0.5)


def test_margin_is_none_without_finance_data(data):
    assert all(brand.margin is None for brand in data.brands)


def test_email_shows_margin_beside_the_value():
    payload = fixture.load()
    finance = build_finance(
        [{"pd_deal_id": 1, "client_name": "X", "stage": "Delivered",
          "budget_gbp": 100_000, "current_spend_gbp": 40_000}],
        [{"campaign_number": "1", "fee_gbp": 40_000}],
    )
    orgs = payload["orgs"]
    org_id = next(iter(orgs))
    report = build_report(**payload, finance=finance, campaign_deal_orgs={1: org_id})
    priced = [b for b in report.brands if b.margin is not None]
    assert priced, "expected at least one brand to carry a margin"
    assert priced[0].margin == pytest.approx(0.6)
    html = render_email(report, title="t", greeting_name="John", sender_name="Valeria")
    assert "#9a99a5" in html  # grey margin styling next to the black value


def test_campaigns_with_a_dead_deal_id_fall_back_to_the_client_name():
    """~111 campaigns carry a pd_deal_id Pipedrive 404s on; keep their money."""
    from app.brands import Brand
    from app.campaign_finance import name_index

    brands = {71: Brand(key="opera.com", name="Opera Browser", org_ids=[71],
                        names=["Opera Browser"], domains={"opera.com"}, won_deals=3)}
    index = name_index(brands)
    # Both the brand's own spelling and its domain label reach the same brand.
    assert index["operabrowser"] == "opera.com"
    assert index["opera"] == "opera.com"


def test_an_ambiguous_client_name_is_dropped_not_guessed():
    from app.brands import Brand
    from app.campaign_finance import name_index

    brands = {
        1: Brand(key="a.com", name="Acme", org_ids=[1], names=["Acme"], domains={"a.com"}, won_deals=1),
        2: Brand(key="b.com", name="Acme", org_ids=[2], names=["Acme"], domains={"b.com"}, won_deals=1),
    }
    assert "acme" not in name_index(brands)


def test_unmatched_campaigns_are_reported_not_hidden():
    payload = fixture.load()
    finance = build_finance(
        [{"pd_deal_id": 999_999, "client_name": "Ghost Client Ltd", "stage": "Delivered",
          "budget_gbp": 50_000, "current_spend_gbp": 20_000}],
        [],
    )
    report = build_report(**payload, finance=finance, campaign_deal_orgs={})
    assert "Ghost Client Ltd" in report.campaigns_unmatched


# -- forecast margins ------------------------------------------------------
def test_a_running_campaign_uses_planned_spend_and_is_flagged():
    finance = build_finance(
        [{"pd_deal_id": 1, "client_name": "X", "stage": "WIP",
          "budget_gbp": 100_000, "planned_spend": 45_000}],
        [],
    )
    rolled = finance.roll_up({1})
    assert rolled.campaigns == 1
    assert rolled.margin == pytest.approx(0.55)
    assert rolled.is_forecast is True


def test_delivered_campaigns_win_over_a_forecast():
    """A measured margin is never blended with a predicted one."""
    finance = build_finance(
        [{"pd_deal_id": 1, "client_name": "X", "stage": "Delivered",
          "budget_gbp": 100_000, "current_spend_gbp": 60_000},
         {"pd_deal_id": 2, "client_name": "X", "stage": "WIP",
          "budget_gbp": 100_000, "planned_spend": 10_000}],
        [],
    )
    rolled = finance.roll_up({1, 2})
    assert rolled.campaigns == 1
    assert rolled.is_forecast is False
    assert rolled.margin == pytest.approx(0.40)


def test_email_stars_a_forecast_margin_and_explains_it():
    payload = fixture.load()
    finance = build_finance(
        [{"pd_deal_id": 1, "client_name": "X", "stage": "WIP",
          "budget_gbp": 100_000, "planned_spend": 45_000}],
        [],
    )
    org_id = next(iter(payload["orgs"]))
    report = build_report(**payload, finance=finance, campaign_deal_orgs={1: org_id})
    assert report.has_forecast_margin is True
    html = render_email(report, title="t", greeting_name="John", sender_name="Valeria")
    assert "55%*" in html
    assert "forecast cost - the campaign is still running" in html


def test_the_margin_note_is_absent_when_nothing_is_priced(data):
    html = render_email(data, title="t", greeting_name="John", sender_name="Valeria")
    assert "gross margin on delivered campaigns" not in html


# -- client aliases --------------------------------------------------------
from app.client_aliases import alias_map, parse_aliases


def test_aliases_parse_from_either_form():
    assert parse_aliases("Match.com LLC = Match Group") == {"matchcom": "match"}
    assert parse_aliases('{"Match.com LLC": "Match Group"}') == {"matchcom": "match"}
    assert parse_aliases("A = B; C = D") == {"a": "b", "c": "d"}


def test_bad_alias_configuration_is_ignored_not_fatal():
    assert parse_aliases("") == {}
    assert parse_aliases("no equals sign here") == {}
    assert parse_aliases("{not json") == {}
    # A name aliased to itself would be a no-op, so it is dropped.
    assert parse_aliases("Same = Same") == {}


def test_configured_aliases_override_the_defaults():
    assert alias_map()["matchcom"] == "match"
    assert alias_map("Match.com LLC = Something Else")["matchcom"] == "somethingelse"


def test_an_alias_attaches_a_campaign_to_the_named_brand():
    payload = fixture.load()
    brand_name = build_report(**payload).brands[0].name
    finance = build_finance(
        [{"pd_deal_id": 777, "client_name": "Totally Different Ltd", "stage": "Delivered",
          "budget_gbp": 100_000, "current_spend_gbp": 40_000}],
        [],
    )
    # Without an alias the campaign has nowhere to go.
    plain = build_report(**payload, finance=finance, campaign_deal_orgs={}, client_aliases={})
    assert "Totally Different Ltd" in plain.campaigns_unmatched

    aliased = build_report(
        **payload, finance=finance, campaign_deal_orgs={},
        client_aliases=parse_aliases(f"Totally Different Ltd = {brand_name}"),
    )
    assert aliased.campaigns_unmatched == []
    brand = next(b for b in aliased.brands if b.name == brand_name)
    assert brand.margin == pytest.approx(0.6)


# -- year to date beside each person ---------------------------------------
def test_ytd_totals_split_by_owner_manager_and_contact():
    payload = fixture.load()
    owner_id = list(payload["users"])[0]
    won = [{"id": 1, "title": "W1", "org_id": None, "owner_id": owner_id, "person_id": None,
            "value": 60_000, "currency": "GBP", "won_time": "2026-06-01T10:00:00Z"},
           {"id": 2, "title": "W2", "org_id": None, "owner_id": owner_id, "person_id": None,
            "value": 40_000, "currency": "GBP", "won_time": "2026-07-01T10:00:00Z"}]
    report = build_report(**payload, won_deals_payload=won,
                          financial_year_start=_dt.date(2026, 4, 1))
    owner = report.won_ytd[0].account_owner
    assert report.ytd_for_owner(owner) == pytest.approx(100_000)
    assert report.ytd_for_owner("Nobody") == 0
    # A blank manager or contact never becomes a bucket of its own.
    assert report.ytd_for_manager("") == 0
    assert report.ytd_for_contact("") == 0


def test_the_podium_shows_won_year_to_date_beside_the_pipeline():
    payload = fixture.load()
    owner_id = list(payload["users"])[0]
    won = [{"id": 1, "title": "W", "org_id": None, "owner_id": owner_id, "person_id": None,
            "value": 250_000, "currency": "GBP", "won_time": "2026-06-01T10:00:00Z"}]
    report = build_report(**payload, won_deals_payload=won,
                          financial_year_start=_dt.date(2026, 4, 1))
    html = render_email(report, title="t", greeting_name="John", sender_name="Valeria")
    assert "won YTD" in html
    assert compact_gbp(250_000) in html


def test_no_ytd_text_when_nothing_has_been_won(data):
    html = render_email(data, title="t", greeting_name="John", sender_name="Valeria")
    assert "won YTD" not in html


# -- the workbook must carry its own numbers -------------------------------
@pytest.fixture(scope="module")
def cached_workbook(data):
    """The workbook as a reader that does not calculate would see it."""
    return openpyxl.load_workbook(io.BytesIO(build_workbook(data).getvalue()), data_only=True)


def test_computed_cells_store_their_result_not_only_a_formula(cached_workbook, data):
    """iPhone Mail and Quick Look never calculate; they show what is stored.

    Without a cached result every figure read as blank or zero on a phone.
    """
    summary = cached_workbook["Summary"]
    assert summary["B5"].value == data.open_deal_count
    assert summary["D5"].value == pytest.approx(data.pipeline_gbp, abs=0.01)
    assert summary["H5"].value == pytest.approx(data.weighted_gbp, abs=0.01)
    assert summary["J5"].value == pytest.approx(data.new_business_gbp, abs=0.01)


def test_deal_and_brand_rows_store_their_results_too(cached_workbook, data):
    deals = cached_workbook["Open Deals"]
    first = data.deals[0]
    assert deals["I2"].value == pytest.approx(first.value_gbp, abs=0.01)
    assert deals["K2"].value == pytest.approx(first.weighted_gbp, abs=0.01)

    brands = cached_workbook["Brands"]
    top = data.brands[0]
    assert brands["F2"].value == top.open_deals
    assert brands["G2"].value == pytest.approx(top.pipeline_gbp, abs=0.01)
    assert brands["H2"].value == top.furthest_stage


def test_stage_and_owner_tables_store_their_results(cached_workbook, data):
    summary = cached_workbook["Summary"]
    stage = data.by_stage()[0]
    assert summary["C10"].value == stage["deals"]
    assert summary["E10"].value == pytest.approx(stage["value_gbp"], abs=0.01)
    owner = next(r for r in data.by_owner() if r["name"] == data.account_owners[0])
    assert summary["I10"].value == owner["deals"]
    assert summary["J10"].value == pytest.approx(owner["value_gbp"], abs=0.01)


def test_the_formulas_survive_the_cached_values(workbook):
    """Both, so the sheet stays auditable and still reads on a phone."""
    assert str(workbook["Summary"]["D5"].value).startswith("=SUM(")
    assert str(workbook["Open Deals"]["I2"].value).startswith("=")
